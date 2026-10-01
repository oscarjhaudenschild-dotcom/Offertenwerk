#!/usr/bin/env python3
"""
Build the scoring sheet (Excel) for the answer-quality sample.

    python3 build_scoring_sheet.py

Reads measurements/generation_sample.csv and writes
measurements/bewertungsbogen.xlsx with one row per question. Every label is
assigned by hand. The sheet only proposes a label in two cases, both marked
"vorausgefüllt": the unanswerable questions, and answers whose figures match
the gold answer exactly (extras still need checking).
"""

import argparse
import csv
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from answer_scoring import REFUSAL_MARKER, score_answer

LABELS = ["U-OK", "U-ERF", "R", "Q", "A", "S-ABL", "S-F", "S-R"]
COLUMNS = [
    ("Frage-ID", 11), ("Kategorie", 15), ("Frage", 42), ("Gold-Antwort", 24),
    ("Treffer", 9), ("Relevante abgerufene Stelle(n)", 70), ("Erzeugte Antwort", 55),
    ("Vorschlag", 26), ("Label", 10), ("Begründung", 40),
]
LABEL_COLUMN = "I"
WRAP = Alignment(wrap_text=True, vertical="top")
HINT = PatternFill("solid", fgColor="FFF2CC")
HEADER = PatternFill("solid", fgColor="1F3864")


def relevant_chunks(row: dict, chunks: dict) -> list:
    """The retrieved chunks that hold the gold answer, if any."""
    expected = [i for i in row["expected_chunk_ids"].split("|") if i in chunks]
    if expected:
        return expected
    wording = row["expected_answer"].split(" -")[0]
    return [i for i, text in chunks.items() if wording and wording in text]


def proposal(row: dict) -> str:
    answer = row["actual_answer"]
    if row["unanswerable"] == "True":
        label = "U-OK" if REFUSAL_MARKER in answer.lower() else "U-ERF"
        return f"{label} (vorausgefüllt)"
    if score_answer(answer, row["expected_answer"])["verdict"] == "correct":
        label = "R" if row["retrieval_hit"] == "True" else "S-R"
        return f"{label} (vorausgefüllt, Zusätze prüfen)"
    return ""


def scoring_rows(results: list) -> list:
    rows = []
    for r in results:
        chunks = json.loads(r["retrieved_chunk_texts"])
        unanswerable = r["unanswerable"] == "True"
        found = relevant_chunks(r, chunks)
        if unanswerable:
            passage = "(unbeantwortbare Frage: keine passende Stelle vorhanden)"
        elif found:
            passage = "\n\n".join(f"[{i}]\n{chunks[i]}" for i in found)
        else:
            passage = "(keine passende Stelle unter den abgerufenen Chunks)"
        rows.append([
            r["case_id"], r["category"], r["query"], r["expected_answer"],
            "–" if unanswerable else ("ja" if r["retrieval_hit"] == "True" else "nein"),
            passage, r["actual_answer"], proposal(r), None, None,
        ])
    return rows


def add_scoring_sheet(wb: Workbook, results: list) -> None:
    ws = wb.active
    ws.title = "Bewertung"
    ws.append([name for name, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
    for index, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    for row in scoring_rows(results):
        ws.append(row)
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = WRAP
        if line[7].value:
            line[7].fill = HINT
    dropdown = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
    ws.add_data_validation(dropdown)
    dropdown.add(f"{LABEL_COLUMN}2:{LABEL_COLUMN}{ws.max_row}")
    ws.freeze_panes = "C2"


def add_evaluation_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Auswertung")
    count = lambda label: f'COUNTIF(Bewertung!${LABEL_COLUMN}:${LABEL_COLUMN},"{label}")'
    ws.append(["Label", "Anzahl"])
    for label in LABELS:
        ws.append([label, f"={count(label)}"])
    ws.append([])
    ws.append(["Quellentreue = R / (R + Q + A)", f"=IFERROR({count('R')}/({count('R')}+{count('Q')}+{count('A')}),\"\")"])
    ws.append(["Ehrlichkeit = U-OK / 5", f"={count('U-OK')}/5"])
    ws.append(["Falsche Antworten vom Suchschritt (S-F)", f"={count('S-F')}"])
    ws.append(["Falsche Antworten vom Antwortschritt (Q + A)", f"={count('Q')}+{count('A')}"])
    ws.append([])
    ws.append(["In der Arbeit als «12 von 42» angeben, nicht als Prozent mit Nachkommastellen."])
    ws.column_dimensions["A"].width = 48


def add_settings_sheet(wb: Workbook, settings: dict) -> None:
    ws = wb.create_sheet("Einstellungen")
    for key, value in settings.items():
        ws.append([key, value])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110
    for line in ws.iter_rows():
        line[1].alignment = WRAP


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results", default="measurements/generation_sample.csv")
    parser.add_argument("--settings", default="measurements/generation_sample_settings.json")
    parser.add_argument("--out", default="measurements/bewertungsbogen.xlsx")
    args = parser.parse_args()

    with open(args.results, encoding="utf-8") as f:
        results = list(csv.DictReader(f))
    with open(args.settings, encoding="utf-8") as f:
        settings = json.load(f)

    wb = Workbook()
    add_scoring_sheet(wb, results)
    add_evaluation_sheet(wb)
    add_settings_sheet(wb, settings)
    wb.save(args.out)
    print(f"{len(results)} Fragen -> {args.out}")


if __name__ == "__main__":
    main()
